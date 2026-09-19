"""Write result tables as XLSX (openpyxl) or CSV depending on the file suffix."""

from __future__ import annotations

import csv
from pathlib import Path
from typing import Iterable, Mapping, Sequence


Row = Mapping[str, object]


def write_table(
    rows: Sequence[Row],
    columns: Sequence[str],
    path: str | Path,
    *,
    sheet: str = "summary",
    extra_sheets: Mapping[str, tuple[Sequence[str], Sequence[Sequence[object]]]] | None = None,
) -> Path:
    """Write ``rows`` (dicts) with the given column order to ``path``.

    ``.xlsx`` -> Excel workbook (needs openpyxl); anything else -> CSV.
    ``extra_sheets`` maps a sheet name to ``(header, rows)`` and is used only
    for XLSX output. Returns the written path.
    """

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    if path.suffix.lower() == ".xlsx":
        _write_xlsx(rows, columns, path, sheet=sheet, extra_sheets=extra_sheets or {})
    else:
        _write_csv(rows, columns, path)
    return path


def _csv_value(value: object) -> object:
    if value is None:
        return ""
    if isinstance(value, float):
        return f"{value:.6f}"
    return value


def _write_csv(rows: Iterable[Row], columns: Sequence[str], path: Path) -> None:
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(columns))
        writer.writeheader()
        for row in rows:
            writer.writerow({key: _csv_value(row.get(key)) for key in columns})


def _write_xlsx(
    rows: Sequence[Row],
    columns: Sequence[str],
    path: Path,
    *,
    sheet: str,
    extra_sheets: Mapping[str, tuple[Sequence[str], Sequence[Sequence[object]]]],
) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
        from openpyxl.utils import get_column_letter
    except ImportError as exc:  # pragma: no cover - depends on the environment
        raise SystemExit(
            f"Writing {path.name} needs openpyxl: pip install openpyxl "
            '(or pip install -e "VaspTools[xlsx]"). Use a .csv output path otherwise.'
        ) from exc

    book = Workbook()
    main = book.active
    main.title = sheet
    _fill_sheet(main, columns, [[row.get(key) for key in columns] for row in rows], Font, get_column_letter)
    for name, (header, data) in extra_sheets.items():
        _fill_sheet(book.create_sheet(name), header, data, Font, get_column_letter)
    book.save(path)


def _fill_sheet(sheet, header: Sequence[str], data: Sequence[Sequence[object]], Font, get_column_letter) -> None:
    sheet.append(list(header))
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for values in data:
        sheet.append([_xlsx_value(value) for value in values])
    sheet.freeze_panes = "B2"
    if data:
        sheet.auto_filter.ref = sheet.dimensions
    for index, key in enumerate(header, start=1):
        width = max([len(str(key))] + [len(str(values[index - 1])) for values in data if values[index - 1] is not None])
        sheet.column_dimensions[get_column_letter(index)].width = min(max(10, width + 2), 60)
        if key.startswith(("energy", "delta_energy")):
            number_format = "0.000000"
        elif key.startswith(("volume", "density", "delta_volume", "runtime")):
            number_format = "0.000"
        else:
            continue
        for column in sheet.iter_cols(min_col=index, max_col=index, min_row=2):
            for cell in column:
                cell.number_format = number_format


def _xlsx_value(value: object) -> object:
    # openpyxl accepts scalars only; keep booleans/numbers, stringify the rest.
    if value is None or isinstance(value, (bool, int, float, str)):
        return value
    return str(value)
