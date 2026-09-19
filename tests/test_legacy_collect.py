from __future__ import annotations

import csv
import shutil
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))

from VaspTools.structures import scale_structure_to_volume, write_poscar

from test_molecules import two_water_conformations

SCRIPT = Path(__file__).resolve().parents[1] / "scripts" / "collect_legacy_relax.py"
DONE = (
    "free  energy   TOTEN  = {a} eV\nfree  energy   TOTEN  = {b} eV\n reached required accuracy\n"
    " General timing and accounting informations for this job:\n Elapsed time (sec): 10.0\n"
)
RUNNING = "free  energy   TOTEN  = {a} eV\nfree  energy   TOTEN  = {b} eV\n"


def make_legacy_run(root: Path, idx: str, *, finished_steps: int, running_tail: bool) -> None:
    calc = root / f"calc_{idx}"
    work = calc / f"calc_{idx}"
    work.mkdir(parents=True)
    structure = two_water_conformations()
    write_poscar(structure, calc / f"POSCAR_{idx}")
    for n in range(1, 4):
        (work / f"step_{n}_initial").mkdir()
        (work / f"step_{n}_initial" / "POSCAR").write_text("placeholder\n")
        if n <= finished_steps:
            final = work / f"step_{n}_final"
            final.mkdir()
            (final / "OUTCAR").write_text(DONE.format(a=-10.0 - n, b=-10.5 - n))
            write_poscar(scale_structure_to_volume(structure, volume_factor=1 - 0.1 * n), final / "CONTCAR")
    last = work / f"step_{finished_steps}_final"
    if running_tail:
        (work / "OUTCAR").write_text(RUNNING.format(a=-13.0, b=-13.2))
        write_poscar(scale_structure_to_volume(structure, volume_factor=0.5), work / "CONTCAR")
    else:
        shutil.copy(last / "OUTCAR", work / "OUTCAR")
        shutil.copy(last / "CONTCAR", work / "CONTCAR")


class LegacyCollectTests(unittest.TestCase):
    def test_collects_table_and_last_contcars(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            run2 = root / "run2"
            make_legacy_run(run2, "00046", finished_steps=3, running_tail=False)
            make_legacy_run(run2, "00665", finished_steps=2, running_tail=True)
            xlsx = root / "summary.xlsx"

            subprocess.run(
                [sys.executable, str(SCRIPT), "--runs-dir", str(run2), "--xlsx", str(xlsx)],
                check=True,
                capture_output=True,
            )

            with xlsx.with_suffix(".csv").open(newline="", encoding="utf-8") as handle:
                rows = {row["idx"]: row for row in csv.DictReader(handle)}
            done = rows["00046"]
            self.assertEqual(done["status"], "completed")
            self.assertEqual(done["steps_finished"], "3")
            self.assertEqual(done["energy_initial_eV"], "-11.0")
            self.assertEqual(done["energy_final_eV"], "-13.5")
            self.assertEqual(done["energy_step2_eV"], "-12.5")
            self.assertEqual(done["n_molecules_initial"], "2")
            self.assertEqual(done["n_molecules_final"], "2")
            self.assertAlmostEqual(float(done["delta_volume_percent"]), -30.0, places=6)
            self.assertGreater(float(done["density_final_g_cm3"]), float(done["density_initial_g_cm3"]))
            self.assertTrue(done["relaxed_path"].endswith("00046_POSCAR"))

            tail = rows["00665"]
            self.assertEqual(tail["status"], "running_or_killed")
            self.assertEqual(tail["energy_final_eV"], "-12.5")  # from step_2_final, not the running dir
            self.assertEqual(tail["relaxed_path"], "")

            relaxed = root / "run2_relaxed"
            self.assertEqual(sorted(path.name for path in relaxed.iterdir()), ["00046_POSCAR"])
            self.assertEqual(
                (relaxed / "00046_POSCAR").read_bytes(),
                (run2 / "calc_00046" / "calc_00046" / "step_3_final" / "CONTCAR").read_bytes(),
            )

            try:
                import openpyxl
            except ImportError:
                return
            book = openpyxl.load_workbook(xlsx)
            self.assertEqual(book.sheetnames, ["relaxations", "steps"])
            self.assertEqual(book["relaxations"].max_row, 3)

    def test_include_unfinished_exports_running_contcar(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            make_legacy_run(root / "run2", "00665", finished_steps=2, running_tail=True)
            subprocess.run(
                [
                    sys.executable, str(SCRIPT), "--runs-dir", str(root / "run2"),
                    "--xlsx", str(root / "s.xlsx"), "--relaxed-dir", str(root / "out"), "--include-unfinished",
                ],
                check=True,
                capture_output=True,
            )
            self.assertTrue((root / "out" / "00665_POSCAR").exists())


if __name__ == "__main__":
    unittest.main()
