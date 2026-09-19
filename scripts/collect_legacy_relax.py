#!/usr/bin/env python3
"""Harvest results of the legacy multi-step relaxation pipeline into XLSX + CONTCARs.

Expected layout (one folder per structure, produced by the old job.py driver)::

    run2/
      calc_00046/
        POSCAR_00046                 # input structure
        calc_00046/                  # VASP working directory (last/current step)
          OUTCAR CONTCAR OSZICAR ...
          step_1_initial/  INCAR POSCAR
          step_1_final/    OUTCAR CONTCAR OSZICAR ...
          step_2_initial/ ...
          step_2_final/   ...

For every calc_* folder the script finds the last finished step, reads energies
(initial = first TOTEN of step 1, final = last TOTEN of the last step), counts
molecules, volume, density and space group before/after, writes one XLSX (+ CSV)
and copies the last CONTCAR to <relaxed-dir>/<idx>_POSCAR.

Example:
  python collect_legacy_relax.py --runs-dir /path/run2 \
      --xlsx /path/run2_summary.xlsx --relaxed-dir /path/run2_relaxed
"""

from __future__ import annotations

import argparse
import csv
import re
import shutil
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
PROJECT_PARENT = ROOT.parent
if str(PROJECT_PARENT) not in sys.path:
    sys.path.insert(0, str(PROJECT_PARENT))

from VaspTools.io.results import (
    read_converged,
    read_energy,
    read_finished,
    read_initial_energy,
)
from VaspTools.io.runtime import read_runtime_seconds
from VaspTools.scripts._scan_utils import describe_structure, load_structure


STEP_FINAL_RE = re.compile(r"^step_(\d+)_final$")
STEP_INITIAL_RE = re.compile(r"^step_(\d+)_initial$")

BASE_COLUMNS = (
    "idx",
    "calc_dir",
    "status",
    "converged",
    "steps_finished",
    "last_step",
    "energy_initial_eV",
    "energy_final_eV",
    "delta_energy_eV",
    "energy_final_per_molecule_eV",
    "n_atoms",
    "formula",
    "n_molecules_initial",
    "n_molecules_final",
    "space_group_initial",
    "space_group_final",
    "volume_initial_A3",
    "volume_final_A3",
    "delta_volume_percent",
    "density_initial_g_cm3",
    "density_final_g_cm3",
    "runtime_sec",
    "relaxed_path",
)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Collect legacy step_N_final relaxations into an XLSX table and a folder of CONTCARs."
    )
    parser.add_argument("--runs-dir", required=True, help="Directory with calc_XXXXX folders (e.g. run2).")
    parser.add_argument("--xlsx", required=True, help="Output XLSX path (a CSV with the same stem is written too).")
    parser.add_argument(
        "--relaxed-dir",
        default=None,
        help="Where to copy the last CONTCAR of every run as <idx>_POSCAR (default: <runs-dir>_relaxed).",
    )
    parser.add_argument(
        "--include-unfinished",
        action="store_true",
        help="Also copy CONTCAR of runs whose last step did not finish normally (marked in the table).",
    )
    parser.add_argument("--symprec", type=float, default=0.05, help="Symmetry tolerance in Angstrom (default 0.05).")
    parser.add_argument(
        "--bond-tolerance", type=float, default=1.20, help="Covalent-radius factor for molecule counting (default 1.20)."
    )
    return parser.parse_args()


def structure_index(calc_dir: Path) -> str:
    """'calc_00046' -> '00046'; falls back to the folder name."""

    match = re.search(r"(\d+)$", calc_dir.name)
    return match.group(1) if match else calc_dir.name


def find_workdir(calc_dir: Path) -> Path:
    """The VASP working directory: calc_XXXXX/calc_XXXXX if present, else calc_XXXXX itself."""

    nested = calc_dir / calc_dir.name
    if nested.is_dir():
        return nested
    candidates = [path for path in calc_dir.iterdir() if path.is_dir() and (path / "OUTCAR").exists()]
    return candidates[0] if len(candidates) == 1 else calc_dir


def find_initial_poscar(calc_dir: Path, workdir: Path, idx: str) -> Path | None:
    for candidate in (
        calc_dir / f"POSCAR_{idx}",
        workdir / "step_1_initial" / "POSCAR",
        *sorted(calc_dir.glob("POSCAR_*")),
    ):
        if candidate.is_file() and candidate.stat().st_size > 0:
            return candidate
    return None


def numbered_dirs(workdir: Path, pattern: re.Pattern[str]) -> list[tuple[int, Path]]:
    found = []
    for path in workdir.iterdir():
        match = pattern.match(path.name)
        if match and path.is_dir():
            found.append((int(match.group(1)), path))
    return sorted(found)


def has_outputs(directory: Path) -> bool:
    return (directory / "OUTCAR").exists() or (directory / "OSZICAR").exists()


def collect_calc(calc_dir: Path, args: argparse.Namespace) -> tuple[dict[str, object], dict[int, float | None]]:
    idx = structure_index(calc_dir)
    workdir = find_workdir(calc_dir)
    row: dict[str, object] = {"idx": idx, "calc_dir": str(calc_dir)}

    finals = [(n, path) for n, path in numbered_dirs(workdir, STEP_FINAL_RE) if has_outputs(path)]
    initials = numbered_dirs(workdir, STEP_INITIAL_RE)
    planned_steps = max([n for n, _ in initials] + [n for n, _ in finals] + [0])

    # The working directory holds the step that is running (or was killed) after
    # the last archived step_N_final; treat it as one more, possibly unfinished, step.
    work_step: tuple[int, Path] | None = None
    if has_outputs(workdir):
        next_n = (finals[-1][0] + 1) if finals else 1
        if next_n <= max(planned_steps, next_n) and read_finished(workdir) is False:
            work_step = (next_n, workdir)
        elif not finals:
            work_step = (next_n, workdir)

    step_energies: dict[int, float | None] = {}
    for n, path in finals:
        try:
            step_energies[n] = read_energy(path)
        except (FileNotFoundError, ValueError):
            step_energies[n] = None

    row["steps_finished"] = len(finals)
    if not finals and work_step is None:
        row["status"] = "missing_outputs"
        row["last_step"] = None
    else:
        last_n, last_dir = finals[-1] if finals else work_step  # type: ignore[misc]
        last_finished = read_finished(last_dir)
        if work_step is not None and finals:
            row["status"] = "running_or_killed"
        elif last_finished is False:
            row["status"] = "running_or_killed"
        elif planned_steps and last_n < planned_steps:
            row["status"] = "partial"
        else:
            row["status"] = "completed"
        row["last_step"] = last_n
        row["converged"] = read_converged(last_dir)

        first_dir = finals[0][1] if finals else last_dir
        try:
            row["energy_initial_eV"] = read_initial_energy(first_dir)
        except (FileNotFoundError, ValueError):
            row["energy_initial_eV"] = None
        try:
            row["energy_final_eV"] = read_energy(last_dir)
        except (FileNotFoundError, ValueError):
            row["energy_final_eV"] = None
        if row["energy_initial_eV"] is not None and row["energy_final_eV"] is not None:
            row["delta_energy_eV"] = float(row["energy_final_eV"]) - float(row["energy_initial_eV"])

        runtimes = [read_runtime_seconds(path) for _, path in finals]
        if work_step is not None:
            runtimes.append(read_runtime_seconds(work_step[1]))
        runtimes = [value for value in runtimes if value is not None]
        row["runtime_sec"] = float(sum(runtimes)) if runtimes else None

        contcar = last_dir / "CONTCAR"
        if contcar.is_file() and contcar.stat().st_size > 0:
            row["_contcar"] = contcar

    initial_poscar = find_initial_poscar(calc_dir, workdir, idx)
    if initial_poscar is not None:
        try:
            description = describe_structure(
                load_structure(initial_poscar),
                symprec=args.symprec,
                bond_tolerance_factor=args.bond_tolerance,
            )
            row["n_atoms"] = description["n_atoms"]
            row["formula"] = description["formula"]
            row["n_molecules_initial"] = description["n_molecules"]
            row["space_group_initial"] = description["space_group"]
            row["volume_initial_A3"] = description["volume_A3"]
            row["density_initial_g_cm3"] = description["density_g_cm3"]
        except Exception as exc:  # unreadable POSCAR should not kill the whole harvest
            print(f"{calc_dir.name}: cannot read initial structure {initial_poscar}: {exc}", file=sys.stderr)

    contcar = row.get("_contcar")
    if contcar is not None:
        try:
            description = describe_structure(
                load_structure(contcar),
                symprec=args.symprec,
                bond_tolerance_factor=args.bond_tolerance,
            )
            row["n_molecules_final"] = description["n_molecules"]
            row["space_group_final"] = description["space_group"]
            row["volume_final_A3"] = description["volume_A3"]
            row["density_final_g_cm3"] = description["density_g_cm3"]
            if row.get("volume_initial_A3"):
                row["delta_volume_percent"] = 100.0 * (
                    float(description["volume_A3"]) / float(row["volume_initial_A3"]) - 1.0
                )
            n_mol = description["n_molecules"] or row.get("n_molecules_initial")
            if n_mol and row.get("energy_final_eV") is not None:
                row["energy_final_per_molecule_eV"] = float(row["energy_final_eV"]) / int(n_mol)
        except Exception as exc:
            print(f"{calc_dir.name}: cannot read {contcar}: {exc}", file=sys.stderr)

    return row, step_energies


def export_contcar(row: dict[str, object], relaxed_dir: Path, *, include_unfinished: bool) -> None:
    contcar = row.get("_contcar")
    if contcar is None:
        return
    if row["status"] != "completed" and not include_unfinished:
        return
    relaxed_dir.mkdir(parents=True, exist_ok=True)
    target = relaxed_dir / f"{row['idx']}_POSCAR"
    shutil.copy2(contcar, target)  # verbatim copy: keeps the VASP CONTCAR exactly
    row["relaxed_path"] = str(target)


def write_outputs(rows: list[dict[str, object]], steps: dict[str, dict[int, float | None]], xlsx: Path) -> None:
    max_step = max([n for per_run in steps.values() for n in per_run] + [0])
    step_columns = [f"energy_step{n}_eV" for n in range(1, max_step + 1)]
    columns = list(BASE_COLUMNS[:9]) + step_columns + list(BASE_COLUMNS[9:])
    for row in rows:
        for n in range(1, max_step + 1):
            row[f"energy_step{n}_eV"] = steps.get(str(row["idx"]), {}).get(n)

    xlsx.parent.mkdir(parents=True, exist_ok=True)
    csv_path = xlsx.with_suffix(".csv")
    with csv_path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=columns)
        writer.writeheader()
        for row in rows:
            writer.writerow({key: ("" if row.get(key) is None else row.get(key)) for key in columns})

    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
        from openpyxl.utils import get_column_letter
    except ImportError:
        print(f"openpyxl is not installed; wrote CSV only: {csv_path}", file=sys.stderr)
        return

    book = Workbook()
    sheet = book.active
    sheet.title = "relaxations"
    sheet.append(columns)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in rows:
        sheet.append([row.get(key) for key in columns])
    sheet.freeze_panes = "B2"
    sheet.auto_filter.ref = sheet.dimensions
    for index, key in enumerate(columns, start=1):
        width = max(len(key), *(len(str(row.get(key, ""))) for row in rows)) if rows else len(key)
        sheet.column_dimensions[get_column_letter(index)].width = min(max(10, width + 2), 60)
        if key.startswith(("energy", "delta", "volume", "density", "runtime")):
            for cell in sheet.iter_cols(min_col=index, max_col=index, min_row=2):
                for item in cell:
                    item.number_format = "0.000000" if key.startswith(("energy", "delta_energy")) else "0.000"

    steps_sheet = book.create_sheet("steps")
    steps_sheet.append(["idx", "step", "energy_eV"])
    for cell in steps_sheet[1]:
        cell.font = Font(bold=True)
    for idx, per_run in steps.items():
        for n, energy in sorted(per_run.items()):
            steps_sheet.append([idx, n, energy])
    book.save(xlsx)


def main() -> None:
    args = parse_args()
    runs_dir = Path(args.runs_dir).resolve()
    if not runs_dir.is_dir():
        raise SystemExit(f"Not a directory: {runs_dir}")
    xlsx = Path(args.xlsx).resolve()
    relaxed_dir = (
        Path(args.relaxed_dir).resolve() if args.relaxed_dir else runs_dir.parent / f"{runs_dir.name}_relaxed"
    )

    calc_dirs = sorted(path for path in runs_dir.iterdir() if path.is_dir() and path.name.startswith("calc_"))
    if not calc_dirs:
        raise SystemExit(f"No calc_* folders found in {runs_dir}")

    rows: list[dict[str, object]] = []
    steps: dict[str, dict[int, float | None]] = {}
    for calc_dir in calc_dirs:
        row, step_energies = collect_calc(calc_dir, args)
        export_contcar(row, relaxed_dir, include_unfinished=args.include_unfinished)
        steps[str(row["idx"])] = step_energies
        rows.append(row)
        print(
            f"{calc_dir.name}: {row['status']}, steps={row.get('steps_finished')}, "
            f"E_final={row.get('energy_final_eV')}, rho={row.get('density_final_g_cm3')}"
        )

    for row in rows:
        row.pop("_contcar", None)
    write_outputs(rows, steps, xlsx)

    completed = sum(1 for row in rows if row["status"] == "completed")
    exported = sum(1 for row in rows if row.get("relaxed_path"))
    print(f"\n{len(rows)} runs ({completed} completed); table: {xlsx} (+ .csv)")
    print(f"{exported} CONTCAR(s) copied to {relaxed_dir} as <idx>_POSCAR")


if __name__ == "__main__":
    main()
